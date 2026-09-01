# -*- coding: utf-8 -*-
"""
quick_plot.py — 一键从 CSV/Excel 生成 Origin 图（Skill 的脚本回退通道）。

示例：
    python quick_plot.py --data data.xlsx --sheet Sheet1 ^
        --x-col 1 --y-cols 2,3 --plot-type line_symbol ^
        --x-title "时间 (s)" --y-title "强度 (a.u.)" ^
        --export fig1.png --save fig1.opju

结束码：0=成功。每一步输出 PASS/FAIL 与真实路径。
"""

import argparse
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from origin_session import OriginSession  # noqa: E402


def read_table(path: Path, sheet: str | None):
    """返回 (header:list[str]|None, rows:list[list])，全部为基本类型。"""
    suf = path.suffix.lower()
    if suf in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("缺少 openpyxl：pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            raw = [r for r in csv.reader(f) if any(c.strip() for c in r)]

    def conv(v):
        if v is None or v == "":
            return ""
        try:
            return float(v)
        except (TypeError, ValueError):
            return str(v)

    rows = [[conv(v) for v in r] for r in raw]
    header = None
    if rows and all(isinstance(v, str) for v in rows[0]):
        header = rows[0]
        rows = rows[1:]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return header, rows


def resolve_cols(header, spec: str, kind: str) -> list[int]:
    """列参数支持 1 起始序号或表头名（逗号分隔）。"""
    out = []
    names = [h.strip() for h in (header or [])]
    for tok in spec.split(","):
        tok = tok.strip()
        if not tok:
            continue
        if tok.isdigit():
            out.append(int(tok))
        elif header and tok in names:
            out.append(names.index(tok) + 1)
        else:
            raise SystemExit(f"{kind} 列无效: {tok}（可用表头: {names}）")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="CSV/Excel → Origin 出版级图形")
    ap.add_argument("--data", required=True, help="CSV/XLSX 数据文件")
    ap.add_argument("--sheet", default=None, help="Excel 工作表名")
    ap.add_argument("--x-col", default="1", help="X 列（序号或表头名）")
    ap.add_argument("--y-cols", required=True,
                    help="Y 列列表，如 '2,3' 或 '浓度,吸光度'；需相邻列")
    ap.add_argument("--plot-type", default="line_symbol",
                    choices=["line", "scatter", "line_symbol", "column"])
    ap.add_argument("--graph-name", default="Fig")
    ap.add_argument("--x-title", default=None)
    ap.add_argument("--y-title", default=None)
    ap.add_argument("--x-min", type=float, default=None)
    ap.add_argument("--x-max", type=float, default=None)
    ap.add_argument("--y-min", type=float, default=None)
    ap.add_argument("--y-max", type=float, default=None)
    ap.add_argument("--log-y", action="store_true")
    ap.add_argument("--color", default=None, help="首条曲线颜色 red/blue/...")
    ap.add_argument("--export", default=None, help="导出图文件 .png/.pdf/.tif/.eps")
    ap.add_argument("--save", default=None, help="保存工程 .opju/.opj（绝对路径）")
    ap.add_argument("--hidden", action="store_true", help="后台运行 Origin")
    args = ap.parse_args()

    steps = []

    def step(name, ok, detail=""):
        steps.append(ok)
        print(("[PASS] " if ok else "[FAIL] ") + name +
              (f" — {detail}" if detail else ""))

    data_path = Path(args.data).resolve()
    if not data_path.exists():
        print(f"[FAIL] 数据文件不存在: {data_path}")
        return 1

    try:
        header, rows = read_table(data_path, args.sheet)
        x_cols = resolve_cols(header, args.x_col, "X")
        y_cols = resolve_cols(header, args.y_cols, "Y")
        x_col = x_cols[0]
        step(f"读取数据 {data_path.name}", True,
             f"{len(rows)} 行 × {len(rows[0]) if rows else 0} 列")
    except SystemExit as exc:
        print(f"[FAIL] {exc}")
        return 1

    sess = OriginSession()
    try:
        sess.connect(visible=not args.hidden)
        step("连接 Origin", True)

        book = sess.workbook_new("Data")
        step("新建工作簿", True, book)

        sess.data_put(book, rows, header=header)
        step("写入数据", True)

        graphs = sess.plot_create(book, x_col=x_col, y_cols=y_cols,
                                  plot_type=args.plot_type,
                                  graph_name=args.graph_name)
        step("绘图", True, f"{graphs}")

        g = graphs[0]
        fmts_ok = []
        if args.x_title:
            fmts_ok.append(sess.axis_set(g, "x", title=args.x_title))
        if args.y_title:
            fmts_ok.append(sess.axis_set(g, "y", title=args.y_title))
        axis_cfg = [("x", args.x_min, args.x_max), ("y", args.y_min, args.y_max)]
        for ax, lo, hi in axis_cfg:
            if lo is not None or hi is not None:
                fmts_ok.append(sess.axis_set(g, ax, vmin=lo, vmax=hi))
        if args.log_y:
            fmts_ok.append(sess.axis_set(g, "y", scale="log10"))
        if args.color:
            fmts_ok.append(sess.series_style(g, color=args.color))
        if fmts_ok:
            step("格式化", all(bool(v) for v in fmts_ok))

        if args.export:
            out = Path(args.export).resolve()
            p = sess.graph_export(str(out), graph=g)
            step("导出图形", p.exists(), f"{p} ({p.stat().st_size} bytes)")

        if args.save:
            sp = sess.project_save(args.save)
            step("保存工程", sp.exists(), f"{sp}")
    except Exception as exc:  # noqa: BLE001
        step(f"异常: {exc}", False)
    finally:
        sess.close()

    print("-" * 50)
    ok_all = all(steps)
    print("结果:", "全部成功" if ok_all else "存在失败步骤，见上方 FAIL")
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main())
